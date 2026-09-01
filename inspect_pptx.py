from pptx import Presentation
prs = Presentation('SIH2026-IDEA-Presentation-Format (1).pptx')
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i} ---")
    for j, shape in enumerate(slide.shapes):
        if hasattr(shape, "text"):
            print(f"Shape {j}: {shape.text.replace(chr(10), ' | ')}")
